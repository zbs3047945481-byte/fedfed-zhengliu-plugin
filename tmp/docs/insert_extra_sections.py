from pathlib import Path
from shutil import copy2

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DOC_PATH = Path('/Users/zhubingshuo/毕设资料/2-复制.docx')
BACKUP_PATH = DOC_PATH.with_name('2-复制.backup-before-extra-5.4.docx')
ROOT = Path('/Users/zhubingshuo/plug-and-play-fedfed-fedavg-plugin')
FIG_DIR = ROOT / 'tmp/docs/extra_figures'
FIGS = {
    'diag': FIG_DIR / 'fig_5_12_diagnostic.png',
    'cross': FIG_DIR / 'fig_5_13_cross_optimizer.png',
    'overhead': FIG_DIR / 'fig_5_14_overhead.png',
}

if not BACKUP_PATH.exists():
    copy2(DOC_PATH, BACKUP_PATH)
for path in FIGS.values():
    if not path.exists():
        raise FileNotFoundError(path)

doc = Document(str(DOC_PATH))

BODY_SIZE = 12
BODY_FONT = '宋体'
HEADING_FONT = '黑体'
LINE_PT = 22
FIRST_INDENT_PT = 24


def set_run_font(run, size=BODY_SIZE, name=BODY_FONT, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def set_para_body(p, first=True):
    pf = p.paragraph_format
    pf.line_spacing = Pt(LINE_PT)
    pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(FIRST_INDENT_PT) if first else None
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for r in p.runs:
        set_run_font(r)


def set_heading(p, level):
    pf = p.paragraph_format
    pf.line_spacing = Pt(LINE_PT)
    pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    pf.first_line_indent = None
    pf.space_before = Pt(10 if level == 2 else 6)
    pf.space_after = Pt(6)
    pf.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:
        set_run_font(r, size=14 if level == 2 else 13.5, name=HEADING_FONT, bold=True)


def clear_para(p):
    for r in list(p.runs):
        p._p.remove(r._r)


def add_para_before(ref_el, text, first=True):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r)
    set_para_body(p, first)
    ref_el.addprevious(p._p)
    return p


def add_heading_before(ref_el, text, level=3):
    p = doc.add_paragraph(style=f'Heading {level}')
    r = p.add_run(text)
    set_heading(p, level)
    ref_el.addprevious(p._p)
    return p


def add_caption_before(ref_el, text, table=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = Pt(LINE_PT)
    pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    pf.first_line_indent = None
    pf.space_before = Pt(6 if table else 3)
    pf.space_after = Pt(3)
    pf.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r)
    ref_el.addprevious(p._p)
    return p


def add_image_before(ref_el, path, width=5.45):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = Pt(LINE_PT)
    pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    pf.first_line_indent = None
    pf.space_before = Pt(3)
    pf.space_after = Pt(0)
    pf.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    ref_el.addprevious(p._p)
    return p


def page_break_before(ref_el):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)
    ref_el.addprevious(p._p)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for name, val in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{name}'))
        if node is None:
            node = OxmlElement(f'w:{name}')
            tcMar.append(node)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')


def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        data = kwargs.get(edge)
        element = borders.find(qn(f'w:{edge}'))
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        if data is None:
            element.set(qn('w:val'), 'nil')
            element.set(qn('w:sz'), '0')
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), 'auto')
        else:
            for key, value in data.items():
                element.set(qn(f'w:{key}'), str(value))


def set_table_width(table, width_dxa):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.insert(0, tblW)
    tblW.set(qn('w:type'), 'dxa')
    tblW.set(qn('w:w'), str(width_dxa))
    jc = tblPr.first_child_found_in('w:jc')
    if jc is None:
        jc = OxmlElement('w:jc')
        tblPr.append(jc)
    jc.set(qn('w:val'), 'center')
    layout = tblPr.first_child_found_in('w:tblLayout')
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        tblPr.append(layout)
    layout.set(qn('w:type'), 'fixed')


def set_cell_width(cell, width):
    tcW = cell._tc.get_or_add_tcPr().tcW
    tcW.set(qn('w:type'), 'dxa')
    tcW.set(qn('w:w'), str(width))


def write_cell(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=9):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = Pt(14)
    pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table_before(ref_el, headers, rows, widths, left_cols=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    set_table_width(table, sum(widths))
    left_cols = set(left_cols or [])
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_width(cell, widths[i])
        set_cell_margins(cell)
        write_cell(cell, h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cell = cells[i]
            set_cell_width(cell, widths[i])
            set_cell_margins(cell)
            write_cell(cell, str(v), align=WD_ALIGN_PARAGRAPH.LEFT if i in left_cols else WD_ALIGN_PARAGRAPH.CENTER)
    line = {'val': 'single', 'sz': '8', 'space': '0', 'color': '000000'}
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_border(cell)
            if ri == 0:
                set_cell_border(cell, top=line, bottom=line)
            elif ri == len(table.rows) - 1:
                set_cell_border(cell, bottom=line)
    ref_el.addprevious(table._tbl)
    return table


def subscript_terms(paragraph):
    terms = [('λfd', ('λ', 'fd')), ('λrec', ('λ', 'rec')), ('λx', ('λ', 'x')),
             ('Rd', ('R', 'd')), ('Ed', ('E', 'd')), ('xs', ('x', 's')),
             ('xr', ('x', 'r')), ('xp', ('x', 'p'))]
    text = paragraph.text
    if not any(raw in text for raw, _ in terms):
        return
    clear_para(paragraph)
    i = 0
    while i < len(text):
        match = None
        for raw, pair in terms:
            if text.startswith(raw, i):
                match = (raw, pair)
                break
        if match:
            raw, (base, sub) = match
            r = paragraph.add_run(base)
            set_run_font(r)
            s = paragraph.add_run(sub)
            set_run_font(s)
            s.font.subscript = True
            i += len(raw)
        else:
            r = paragraph.add_run(text[i])
            set_run_font(r)
            i += 1
    set_para_body(paragraph, first=paragraph.paragraph_format.first_line_indent is not None)


def para_texts():
    return [p.text.strip() for p in doc.paragraphs]


texts = para_texts()
start = None
end = None
for p in doc.paragraphs:
    t = p.text.strip()
    if t == '5.4.7 特征蒸馏机制诊断':
        start = p
    if start is not None and t == '5.4.10 结果分析':
        end = p
        break
if start is not None and end is not None:
    body = doc._body._element
    children = list(body)
    si = children.index(start._p)
    ei = children.index(end._p)
    for el in children[si:ei]:
        body.remove(el)

target = None
for p in doc.paragraphs:
    if p.text.strip() == '5.4.7 结果分析':
        clear_para(p)
        r = p.add_run('5.4.10 结果分析')
        set_heading(p, 3)
        target = p
        break
    if p.text.strip() == '5.4.10 结果分析':
        target = p
        break
if target is None:
    raise RuntimeError('Could not locate result analysis heading.')

ref = target._p

add_heading_before(ref, '5.4.7 特征蒸馏机制诊断', 3)
add_para_before(ref, '为进一步解释 FedFed 插件的性能来源，本文对特征蒸馏阶段得到的不同输入形式进行分类诊断。该实验不直接比较联邦训练最终精度，而是检验原始图像、性能敏感特征、性能鲁棒特征以及带噪敏感特征各自保留的任务信息。')
add_para_before(ref, '实验设置六类分类代理任务。其中，x→x 表示使用原始图像训练并在原始图像上测试；xs→xs 表示使用性能敏感特征训练并在性能敏感特征上测试；xr→xr 表示使用性能鲁棒特征训练并在性能鲁棒特征上测试；xs→x 表示使用性能敏感特征训练并迁移到原始图像测试；x+xs→x 表示同时使用原始图像和敏感特征训练，并在原始图像上测试；x+xp→x 表示使用原始图像和带噪敏感特征训练，并在原始图像上测试。')
page_break_before(ref)
add_caption_before(ref, '表 5-8 特征蒸馏机制诊断结果', table=True)
add_table_before(ref, ['代理任务', '输入含义', 'BestAcc', 'FinalAcc', 'BestEpoch'], [
    ['x→x', '原始图像分类', '69.55%', '69.55%', '10'],
    ['xs→xs', '敏感特征分类', '74.00%', '74.00%', '10'],
    ['xr→xr', '鲁棒特征分类', '54.23%', '53.53%', '7'],
    ['xs→x', '敏感特征到原图迁移', '14.29%', '14.29%', '10'],
    ['x+xs→x', '原图与敏感特征联合训练', '73.58%', '71.29%', '8'],
    ['x+xp→x', '原图与带噪敏感特征联合训练', '71.28%', '71.22%', '7'],
], [1100, 3000, 1200, 1200, 1200], left_cols={1})
add_para_before(ref, '表中，BestAcc 表示代理分类器训练过程中的最高测试准确率，FinalAcc 表示训练结束时的测试准确率，BestEpoch 表示取得最高准确率的训练轮次。xs 表示性能敏感特征，xr 表示性能鲁棒特征，xp 表示加入扰动后的性能敏感特征。')
add_para_before(ref, '由表 5-8 可知，xs→xs 的最高准确率达到 74.00%，高于原始图像代理任务 x→x 的 69.55%，说明性能敏感特征中保留了较强的分类判别信息。相比之下，xr→xr 的最高准确率为 54.23%，明显低于 xs→xs，说明性能鲁棒特征中的分类信息相对较弱。该结果符合 FedFed 的特征蒸馏目标：敏感特征主要承载与分类任务相关的信息，鲁棒特征更多保留非关键的隐私信息。')
add_para_before(ref, '同时，x+xs→x 的最高准确率达到 73.58%，高于单独使用原始图像的 x→x，说明敏感特征可以作为有效的辅助训练信息。加入扰动后的 x+xp→x 仍取得 71.28% 的最高准确率，表明扰动后的敏感特征仍然保留任务可用性。结合前文隐私验证结果可知，带噪敏感特征在保持分类辅助作用的同时，可以降低直接反演原始图像的风险。')
add_image_before(ref, FIGS['diag'])
add_caption_before(ref, '图 5-12 不同特征输入形式下的分类诊断结果')
add_para_before(ref, '图 5-12 对比了原始图像、敏感特征、鲁棒特征和带噪敏感特征的分类代理性能。结果显示，敏感特征具有较强分类能力，鲁棒特征分类性较弱，说明 FedFed 插件的特征蒸馏过程能够将更多分类相关信息集中到可共享特征中。')

add_heading_before(ref, '5.4.8 跨联邦优化器可插拔验证', 3)
add_para_before(ref, '前述实验主要基于 FedAvg 主流程验证 FedFed 插件的有效性。为进一步验证插件设计的可插拔性，本文将 FedFed 插件接入 FedProx、SCAFFOLD、FedAvgM 和 FedNova 四种联邦优化器，比较开启插件前后的分类性能。该实验用于验证插件接口能否在不同联邦优化主流程中复用，并分析插件效果与底层优化器之间的关系。')
page_break_before(ref)
add_caption_before(ref, '表 5-9 不同联邦优化器下的插件接入结果', table=True)
add_table_before(ref, ['基础优化器', '插件状态', 'BestAcc', 'FinalAcc', 'BestRound', 'FinalRound', 'Gain'], [
    ['FedProx', '关闭', '59.91%', '58.65%', '133', '168', '-'],
    ['FedProx', '开启', '90.49%', '90.25%', '279', '300', '+30.58%'],
    ['SCAFFOLD', '关闭', '51.68%', '40.33%', '161', '196', '-'],
    ['SCAFFOLD', '开启', '56.27%', '53.54%', '172', '207', '+4.59%'],
    ['FedAvgM', '关闭', '49.52%', '43.87%', '140', '175', '-'],
    ['FedAvgM', '开启', '80.20%', '70.86%', '227', '262', '+30.68%'],
    ['FedNova', '关闭', '47.04%', '42.73%', '64', '120', '-'],
    ['FedNova', '开启', '70.77%', '62.87%', '172', '207', '+23.73%'],
], [1500, 1000, 1150, 1150, 1250, 1300, 1100], left_cols={0})
add_para_before(ref, '表中，基础优化器表示联邦训练使用的主优化方法，插件状态表示是否启用 FedFed 插件，Gain 表示同一基础优化器下开启插件后相对于关闭插件的 BestAcc 提升。')
add_para_before(ref, '由表 5-9 可知，FedFed 插件能够接入多种联邦优化器，并在 FedProx、FedAvgM 和 FedNova 上带来明显性能提升。其中，FedProx 上最高准确率由 59.91% 提升到 90.49%，FedAvgM 上由 49.52% 提升到 80.20%，FedNova 上由 47.04% 提升到 70.77%。该结果说明，本文插件接口不依赖单一 FedAvg 实现，而是可以作为辅助训练模块接入不同联邦优化流程。')
add_para_before(ref, 'SCAFFOLD 上的提升相对有限，BestAcc 由 51.68% 提升到 56.27%。SCAFFOLD 通过服务器端和客户端控制变量校正本地梯度方向，其核心目标是降低 Non-IID 场景下的 client drift；FedFed 插件则通过共享敏感特征改变客户端本地训练数据分布，为本地分类目标提供额外信息。二者都作用于客户端更新方向，但作用路径不同：SCAFFOLD 从梯度校正角度约束更新，FedFed 从数据补充角度改变本地训练信号。当控制变量估计和共享特征辅助目标同时作用时，客户端更新受到两类校正信号共同影响，插件增益被削弱。因此，跨算法实验表明，FedFed 插件具备跨联邦优化器接入能力，但性能收益会随底层优化器机制和训练配置变化。')
add_image_before(ref, FIGS['cross'])
add_caption_before(ref, '图 5-13 不同联邦优化器下的插件性能对比')
add_para_before(ref, '图 5-13 展示了不同基础优化器开启和关闭 FedFed 插件后的最高准确率。可以看到，插件在多数优化器上带来性能提升，但提升幅度并不完全一致，说明插件具有可插拔性，同时其效果与底层联邦优化策略存在耦合关系。')

add_heading_before(ref, '5.4.9 训练开销分析', 3)
add_para_before(ref, 'FedFed 插件通过特征蒸馏和共享特征辅助训练提升模型性能，但也会带来额外训练开销。为客观评价插件的工程代价，本文统计不同设置下无插件基线和 FedFed 插件的训练时间，并计算相对耗时倍率。')
add_caption_before(ref, '表 5-10 FedFed 插件训练开销对比', table=True)
add_table_before(ref, ['实验设置', '方法', 'DistillRounds', '共享池设置', 'BestAcc', 'Duration', '相对倍率'], [
    ['α=0.1, E=1', '无插件基线', '0', '无', '63.89%', '44.11 min', '1.00×'],
    ['α=0.1, E=1', 'FedFed 插件', '15', '不设额外限制', '88.95%', '163.32 min', '3.70×'],
    ['α=0.1, E=5', '无插件基线', '0', '无', '59.60%', '116.03 min', '1.00×'],
    ['α=0.1, E=5', 'FedFed 插件', '15', '不设额外限制', '91.90%', '600.50 min', '5.18×'],
], [1450, 1350, 1400, 1550, 1050, 1250, 1050], left_cols={0, 1, 3})
add_para_before(ref, '表中，DistillRounds 表示特征蒸馏阶段的通信轮数，Duration 表示完整实验运行时间，相对倍率表示同一实验设置下 FedFed 插件方法相对于无插件基线的耗时倍数。共享池设置为“不设额外限制”表示实验中未额外限制共享特征池总容量和单类容量。')
add_para_before(ref, '由表 5-10 可知，在 α=0.1, E=1 的主实验设置下，FedFed 插件将最高准确率从 63.89% 提升到 88.95%，训练时间由 44.11 min 增加到 163.32 min，约为无插件基线的 3.70 倍。在 E=5 的设置下，FedFed 插件最高准确率达到 91.90%，训练时间增加到 600.50 min，约为无插件基线的 5.18 倍。')
add_para_before(ref, '该结果说明，FedFed 插件的性能提升伴随额外计算成本。其开销主要来自两部分：一是正式训练前的特征蒸馏阶段；二是正式训练阶段对共享敏感特征的额外训练。当本地训练轮数增大时，共享特征辅助训练的计算量进一步增加。')
add_para_before(ref, '因此，FedFed 插件适合用于数据异构较强、精度收益需求较高的场景。在实际部署中，需要结合计算资源约束选择蒸馏轮数、共享池规模和本地训练轮数，以在模型性能和训练开销之间取得平衡。')
add_image_before(ref, FIGS['overhead'])
add_caption_before(ref, '图 5-14 FedFed 插件训练时间开销对比')
add_para_before(ref, '图 5-14 展示了无插件基线和 FedFed 插件在 E=1、E=5 两种设置下的训练时间。FedFed 插件带来明显精度提升的同时，也引入额外计算开销，体现了性能收益与工程成本之间的权衡关系。')

# Normalize all 5.4 body paragraphs inserted in earlier turns to the reference body format.
caption_texts = {
    '图 5-5 强异构场景下主性能对比',
    '图 5-6 强异构场景下的收敛曲线',
    '图 5-7 不同数据异构强度下的性能对比',
    '图 5-8 不同数据异构强度下的插件增益',
    '图 5-9 本地训练强度增大时的性能对比',
    '图 5-10 不同共享特征池规模下的性能对比',
    '图 5-11 不同特征蒸馏训练量下的性能对比',
    '图 5-12 不同特征输入形式下的分类诊断结果',
    '图 5-13 不同联邦优化器下的插件性能对比',
    '图 5-14 FedFed 插件训练时间开销对比',
}
in_54 = False
for p in doc.paragraphs:
    t = p.text.strip()
    if t == '5.4 FedFed 插件有效性验证':
        in_54 = True
    elif t == '结  论':
        in_54 = False
    if not in_54 or not t:
        continue
    if t.startswith('5.4 '):
        set_heading(p, 2)
    elif t.startswith('5.4.'):
        set_heading(p, 3)
    elif t.startswith('表 ') or t in caption_texts:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.line_spacing = Pt(LINE_PT)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        for r in p.runs:
            set_run_font(r)
    else:
        set_para_body(p, first=True)
        subscript_terms(p)

settings = doc.settings._element
update = settings.find(qn('w:updateFields'))
if update is None:
    update = OxmlElement('w:updateFields')
    settings.append(update)
update.set(qn('w:val'), 'true')

doc.save(str(DOC_PATH))
print(DOC_PATH)
print(BACKUP_PATH)

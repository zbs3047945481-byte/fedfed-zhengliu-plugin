from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
DOC='/Users/zhubingshuo/毕设资料/2-复制.docx'
d=Document(DOC)
# Replace longer labels first to avoid 图 5-10 -> 图 6-20 type errors.
mapping={
'表 5-10':'表 6-5','图 5-14':'图 6-7','图 5-13':'图 6-6','图 5-12':'图 6-1','表 5-9':'表 6-4','表 5-8':'表 6-1',
'图 5-11':'图 5-7','图 5-10':'图 5-6','表 5-7':'表 5-5','表 5-6':'表 5-4','图 5-9':'图 5-5','表 5-5':'表 5-3','图 5-8':'图 5-4','图 5-7':'图 5-3','表 5-4':'表 5-2','图 5-6':'图 5-2','图 5-5':'图 5-1','表 5-3':'表 5-1',
'图 5-4':'图 6-5','表 5-2':'表 6-3','图 5-3':'图 6-4','图 5-2':'图 6-3','表 5-1':'表 6-2','图 5-1':'图 6-2',
}
items=sorted(mapping.items(), key=lambda kv: len(kv[0]), reverse=True)
for p in d.paragraphs:
    for r in p.runs:
        txt=r.text
        for old,new in items:
            txt=txt.replace(old,new)
        r.text=txt
    # fallback for split runs
    txt=p.text
    if any(old in txt for old,_ in items):
        for old,new in items:
            txt=txt.replace(old,new)
        for r in list(p.runs): p._p.remove(r._r)
        r=p.add_run(txt)
        r.font.name='宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(12)
    t=p.text.strip()
    if t.startswith('图 ') or t.startswith('表 '):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=None
        p.paragraph_format.line_spacing=Pt(22); p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
        p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(3)
        for r in p.runs:
            r.font.name='宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(12)

d.save(DOC)

from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

DOC_PATH = Path('/Users/zhubingshuo/毕设资料/2-复制.docx')
doc = Document(str(DOC_PATH))
body = doc._body._element
BODY_FONT='宋体'; HEADING_FONT='黑体'; BODY_SIZE=12; LINE_PT=22; FIRST_INDENT_PT=24

def set_run_font(run, size=BODY_SIZE, name=BODY_FONT, bold=False):
    run.font.name=name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size=Pt(size)
    run.font.bold=bold

def fmt_body(p, first=True):
    pf=p.paragraph_format
    pf.line_spacing=Pt(LINE_PT); pf.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
    pf.space_before=Pt(0); pf.space_after=Pt(0)
    pf.first_line_indent=Pt(FIRST_INDENT_PT) if first else None
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for r in p.runs: set_run_font(r)

def fmt_heading(p, level):
    pf=p.paragraph_format
    pf.line_spacing=Pt(LINE_PT); pf.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
    pf.first_line_indent=None
    pf.space_before=Pt(22.35 if level==1 else (11.15 if level==2 else 6))
    pf.space_after=Pt(17.85 if level==1 else (11.15 if level==2 else 6))
    pf.keep_with_next=True
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs: set_run_font(r, size=15 if level==1 else (14 if level==2 else 13.5), name=HEADING_FONT, bold=True)

def clear_para(p):
    for r in list(p.runs): p._p.remove(r._r)

def set_text(p, text, level=None):
    clear_para(p)
    p.add_run(text)
    if level:
        try: p.style=f'Heading {level}'
        except Exception: pass
        fmt_heading(p, level)
    else:
        fmt_body(p)

def make_para(text, level=None, first=True):
    p=doc.add_paragraph(style=f'Heading {level}' if level else None)
    p.add_run(text)
    if level: fmt_heading(p, level)
    else: fmt_body(p, first)
    return p

def find_para_exact(text):
    for p in doc.paragraphs:
        if p.text.strip()==text:
            return p
    return None

def elements_between(start_p, end_p):
    children=list(body)
    si=children.index(start_p._p); ei=children.index(end_p._p)
    return children[si:ei]

def insert_elements_before(ref_el, elements):
    for el in elements:
        ref_el.addprevious(el)

def insert_para_before(ref_el, text, level=None, first=True):
    p=make_para(text, level=level, first=first)
    ref_el.addprevious(p._p)
    return p

ch5=find_para_exact('第5章 实验设计与结果分析')
conclusion=find_para_exact('结  论')
if ch5 is None or conclusion is None:
    raise RuntimeError('Cannot locate chapter 5 or conclusion')

p51=find_para_exact('5.1 实验目的'); p52=find_para_exact('5.2 实验环境与基础设置'); p53=find_para_exact('5.3 隐私验证'); p54=find_para_exact('5.4 FedFed 插件有效性验证')
sub_heads=['5.4.1 实验验证思路','5.4.2 主性能对比','5.4.3 数据异构强度分析','5.4.4 本地训练强度分析','5.4.5 共享特征池规模分析','5.4.6 特征蒸馏充分性分析','5.4.7 特征蒸馏机制诊断','5.4.8 跨联邦优化器可插拔验证','5.4.9 训练开销分析','5.4.10 结果分析']
subp=[find_para_exact(h) for h in sub_heads]
if any(x is None for x in [p51,p52,p53,p54,*subp]):
    raise RuntimeError('missing required heading')

sections={}
sections['5.1']=elements_between(p51,p52)
sections['5.2']=elements_between(p52,p53)
sections['privacy']=elements_between(p53,p54)
sections['perf_intro']=elements_between(p54,subp[0])
for i,h in enumerate(sub_heads):
    end=subp[i+1] if i+1<len(subp) else conclusion
    sections[h]=elements_between(subp[i],end)

# Remove only the old chapter-5 title. Other old content will be moved below.
body.remove(ch5._p)
ref=conclusion._p

insert_para_before(ref,'第5章 实验设计与性能评估',level=1)
insert_elements_before(ref,sections['5.1'])
insert_elements_before(ref,sections['5.2'])
insert_para_before(ref,'5.3 评价指标与公平性控制',level=2)
insert_elements_before(ref,sections['perf_intro'][1:])
insert_elements_before(ref,sections['5.4.1 实验验证思路'][1:])
insert_para_before(ref,'为保证实验对比公平，后续实验在相同数据划分、客户端数量、客户端采样比例、主模型结构、优化器设置和服务器聚合规则下进行。除目标变量外，其余实验条件保持一致。',first=True)
insert_para_before(ref,'5.4 FedFed 插件主性能验证',level=2)
for old in ['5.4.2 主性能对比','5.4.3 数据异构强度分析','5.4.4 本地训练强度分析']:
    insert_elements_before(ref,sections[old])
insert_para_before(ref,'5.5 主性能改进历程分析',level=2)
for h in ['5.5.1 初始插件版本','5.5.2 图像空间特征蒸馏版本','5.5.3 共享特征池优化版本','5.5.4 最终插件版本']:
    insert_para_before(ref,h,level=3)
insert_para_before(ref,'5.6 插件关键模块消融实验',level=2)
insert_elements_before(ref,sections['5.4.5 共享特征池规模分析'])
insert_elements_before(ref,sections['5.4.6 特征蒸馏充分性分析'])
for h in ['5.6.3 插件小模块消融一','5.6.4 插件小模块消融二','5.6.5 插件小模块消融三','5.6.6 插件小模块消融四']:
    insert_para_before(ref,h,level=3)
insert_para_before(ref,'5.7 本章小结',level=2)
insert_para_before(ref,'本章围绕 FedFed 插件的性能表现展开实验评估。首先给出实验目的、基础设置和评价指标；随后从主性能、数据异构强度和本地训练强度三个角度验证插件有效性；最后通过性能改进历程和关键模块消融分析不同设计对最终结果的影响。',first=True)

insert_para_before(ref,'第6章 机制验证与工程分析',level=1)
insert_elements_before(ref,sections['5.4.7 特征蒸馏机制诊断'])
insert_elements_before(ref,sections['privacy'])
insert_elements_before(ref,sections['5.4.8 跨联邦优化器可插拔验证'])
insert_elements_before(ref,sections['5.4.9 训练开销分析'])
insert_para_before(ref,'6.5 本章小结',level=2)
insert_para_before(ref,'本章从机制、安全性、可插拔性和工程开销四个角度对 FedFed 插件进行补充分析。实验结果表明，插件能够将分类相关信息集中到性能敏感特征中，并具备跨联邦优化器接入能力；同时，隐私验证和开销分析说明该方法仍需要在隐私风险控制和训练成本之间进行权衡。',first=True)

renames={
 '5.4.2 主性能对比':'5.4.1 主性能对比','5.4.3 数据异构强度分析':'5.4.2 数据异构强度分析','5.4.4 本地训练强度分析':'5.4.3 本地训练强度分析','5.4.5 共享特征池规模分析':'5.6.1 共享特征池规模消融','5.4.6 特征蒸馏充分性分析':'5.6.2 特征蒸馏充分性消融','5.4.7 特征蒸馏机制诊断':'6.1 特征蒸馏机制诊断','5.3 隐私验证':'6.2 隐私验证','5.3.1 隐私验证思路':'6.2.1 隐私验证思路','5.3.2 敏感特征剪裁与噪声设置':'6.2.2 敏感特征剪裁与噪声设置','5.3.3 模型反演攻击验证':'6.2.3 模型反演攻击验证','5.3.4 成员推断攻击验证':'6.2.4 成员推断攻击验证','5.3.5 结果分析':'6.2.5 隐私验证结果分析','5.4.8 跨联邦优化器可插拔验证':'6.3 跨联邦优化器可插拔验证','5.4.9 训练开销分析':'6.4 训练开销分析'}
for p in doc.paragraphs:
    t=p.text.strip()
    if t in renames:
        new=renames[t]; set_text(p,new,level=2 if new.count('.')==1 else 3)
for p in doc.paragraphs:
    t=p.text.strip()
    if t in {'第5章 实验设计与性能评估','第6章 机制验证与工程分析','结  论'}:
        fmt_heading(p,1); p.paragraph_format.page_break_before=True
    elif p.style.name.startswith('Heading') and (t.startswith('5.') or t.startswith('6.')):
        fmt_heading(p,2 if t.split()[0].count('.')==1 else 3)
    elif t.startswith('图 ') or t.startswith('表 '):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent=None
        p.paragraph_format.line_spacing=Pt(LINE_PT); p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
        for r in p.runs: set_run_font(r)

doc.save(str(DOC_PATH))
print('move-based restructure done; inline_shapes=', len(doc.inline_shapes))

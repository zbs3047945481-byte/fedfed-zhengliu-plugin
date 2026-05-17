from pathlib import Path
from shutil import copy2
from copy import deepcopy
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DOC_PATH = Path('/Users/zhubingshuo/毕设资料/2-复制.docx')
BACKUP_PATH = DOC_PATH.with_name('2-复制.backup-before-ch5-ch6-restructure.docx')
if not BACKUP_PATH.exists():
    copy2(DOC_PATH, BACKUP_PATH)

doc = Document(str(DOC_PATH))
BODY_FONT='宋体'; HEADING_FONT='黑体'; BODY_SIZE=12; LINE_PT=22; FIRST_INDENT_PT=24


def set_run_font(run, size=BODY_SIZE, name=BODY_FONT, bold=False):
    run.font.name=name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size=Pt(size)
    run.font.bold=bold


def fmt_body(p, first=True):
    pf=p.paragraph_format
    pf.line_spacing=Pt(LINE_PT); pf.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
    pf.space_before=Pt(0); pf.space_after=Pt(0)
    pf.first_line_indent=Pt(FIRST_INDENT_PT) if first else None
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for r in p.runs: set_run_font(r)


def fmt_heading(p, level):
    pf=p.paragraph_format
    pf.line_spacing=Pt(LINE_PT); pf.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
    pf.first_line_indent=None
    pf.space_before=Pt(22.35 if level==1 else (11.15 if level==2 else 6))
    pf.space_after=Pt(17.85 if level==1 else (11.15 if level==2 else 6))
    pf.keep_with_next=True
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs: set_run_font(r, size=15 if level==1 else (14 if level==2 else 13.5), name=HEADING_FONT, bold=True)


def clear_para(p):
    for r in list(p.runs): p._p.remove(r._r)


def set_text(p, text, level=None):
    clear_para(p)
    r=p.add_run(text)
    if level:
        try: p.style=f'Heading {level}'
        except Exception: pass
        fmt_heading(p, level)
    else:
        fmt_body(p)


def make_para(text, level=None, first=True):
    p=doc.add_paragraph(style=f'Heading {level}' if level else None)
    p.add_run(text)
    if level: fmt_heading(p, level)
    else: fmt_body(p, first)
    return p


def get_idx(el):
    return list(body).index(el)


def find_para_exact(text, start=0):
    for i,p in enumerate(doc.paragraphs[start:], start):
        if p.text.strip()==text:
            return p
    return None


def clone_elements_between(start_p, end_p):
    children=list(body)
    si=children.index(start_p._p); ei=children.index(end_p._p)
    return [deepcopy(el) for el in children[si:ei]]


def remove_between(start_p, end_p):
    children=list(body)
    si=children.index(start_p._p); ei=children.index(end_p._p)
    for el in children[si:ei]: body.remove(el)


def insert_elements_before(ref_el, elements):
    for el in elements:
        ref_el.addprevious(el)


def insert_para_before(ref_el, text, level=None, first=True):
    p=make_para(text, level=level, first=first)
    ref_el.addprevious(p._p)
    return p

body=doc._body._element

# Locate major boundaries.
ch5=find_para_exact('第5章 实验设计与结果分析')
conclusion=find_para_exact('结  论')
if ch5 is None or conclusion is None:
    raise RuntimeError('Cannot locate chapter 5 or conclusion')

# Capture existing sections before deleting chapter 5.
sections={}
heads=['5.1 实验目的','5.2 实验环境与基础设置','5.3 隐私验证','5.4 FedFed 插件有效性验证']
for h in heads:
    p=find_para_exact(h)
    if p is None:
        raise RuntimeError(f'missing {h}')
# Use next headings as ends.
p51=find_para_exact('5.1 实验目的'); p52=find_para_exact('5.2 实验环境与基础设置'); p53=find_para_exact('5.3 隐私验证'); p54=find_para_exact('5.4 FedFed 插件有效性验证')
sections['5.1']=clone_elements_between(p51,p52)
sections['5.2']=clone_elements_between(p52,p53)
sections['privacy']=clone_elements_between(p53,p54)
# sub sections inside 5.4
sub_heads=['5.4.1 实验验证思路','5.4.2 主性能对比','5.4.3 数据异构强度分析','5.4.4 本地训练强度分析','5.4.5 共享特征池规模分析','5.4.6 特征蒸馏充分性分析','5.4.7 特征蒸馏机制诊断','5.4.8 跨联邦优化器可插拔验证','5.4.9 训练开销分析','5.4.10 结果分析']
subp=[find_para_exact(h) for h in sub_heads]
if any(p is None for p in subp):
    raise RuntimeError('missing one or more 5.4 subheads')
# intro content under 5.4 before 5.4.1
sections['perf_intro']=clone_elements_between(p54, subp[0])
for i,h in enumerate(sub_heads):
    end = subp[i+1] if i+1 < len(subp) else conclusion
    sections[h]=clone_elements_between(subp[i], end)

# Delete old chapter 5 content, keeping conclusion onward.
remove_between(ch5, conclusion)
ref=conclusion._p

# Build new chapter 5.
insert_para_before(ref, '第5章 实验设计与性能评估', level=1)
insert_elements_before(ref, sections['5.1'])
insert_elements_before(ref, sections['5.2'])
insert_para_before(ref, '5.3 评价指标与公平性控制', level=2)
# Move old 5.4 intro + validation idea content here, retitle validation idea not included.
# Use body paragraphs cloned from 5.4 intro and 5.4.1 except its heading; then add fairness placeholder text from current explanation if available.
intro=sections['perf_intro']
# skip first heading element from 5.4 section (old heading) and include its paragraphs
insert_elements_before(ref, intro[1:])
val=sections['5.4.1 实验验证思路']
insert_elements_before(ref, val[1:])
insert_para_before(ref, '为保证实验对比公平，后续实验在相同数据划分、客户端数量、客户端采样比例、主模型结构、优化器设置和服务器聚合规则下进行。除目标变量外，其余实验条件保持一致。', first=True)

insert_para_before(ref, '5.4 FedFed 插件主性能验证', level=2)
# Insert 5.4.2-5.4.4 and renumber to 5.4.1-5.4.3
for old,new in [('5.4.2 主性能对比','5.4.1 主性能对比'),('5.4.3 数据异构强度分析','5.4.2 数据异构强度分析'),('5.4.4 本地训练强度分析','5.4.3 本地训练强度分析')]:
    elems=sections[old]
    # change heading text in clone
    first=elems[0]
    insert_elements_before(ref, elems)
    # The inserted paragraph is now before ref; easiest adjust later globally.

insert_para_before(ref, '5.5 主性能改进历程分析', level=2)
insert_para_before(ref, '5.5.1 初始插件版本', level=3)
insert_para_before(ref, '5.5.2 图像空间特征蒸馏版本', level=3)
insert_para_before(ref, '5.5.3 共享特征池优化版本', level=3)
insert_para_before(ref, '5.5.4 最终插件版本', level=3)

insert_para_before(ref, '5.6 插件关键模块消融实验', level=2)
# Existing ablations
for old,new in [('5.4.5 共享特征池规模分析','5.6.1 共享特征池规模消融'),('5.4.6 特征蒸馏充分性分析','5.6.2 特征蒸馏充分性消融')]:
    insert_elements_before(ref, sections[old])
# Empty planned module ablation headings
insert_para_before(ref, '5.6.3 插件小模块消融一', level=3)
insert_para_before(ref, '5.6.4 插件小模块消融二', level=3)
insert_para_before(ref, '5.6.5 插件小模块消融三', level=3)
insert_para_before(ref, '5.6.6 插件小模块消融四', level=3)
insert_para_before(ref, '5.7 本章小结', level=2)
insert_para_before(ref, '本章围绕 FedFed 插件的性能表现展开实验评估。首先给出实验目的、基础设置和评价指标；随后从主性能、数据异构强度和本地训练强度三个角度验证插件有效性；最后通过性能改进历程和关键模块消融分析不同设计对最终结果的影响。', first=True)

# Build new chapter 6.
insert_para_before(ref, '第6章 机制验证与工程分析', level=1)
for old,new in [('5.4.7 特征蒸馏机制诊断','6.1 特征蒸馏机制诊断')]:
    insert_elements_before(ref, sections[old])
# privacy after mechanism
insert_elements_before(ref, sections['privacy'])
for old in ['5.4.8 跨联邦优化器可插拔验证','5.4.9 训练开销分析']:
    insert_elements_before(ref, sections[old])
insert_para_before(ref, '6.5 本章小结', level=2)
insert_para_before(ref, '本章从机制、安全性、可插拔性和工程开销四个角度对 FedFed 插件进行补充分析。实验结果表明，插件能够将分类相关信息集中到性能敏感特征中，并具备跨联邦优化器接入能力；同时，隐私验证和开销分析说明该方法仍需要在隐私风险控制和训练成本之间进行权衡。', first=True)

# Global retitle headings after insertion.
renames={
    '5.4.2 主性能对比':'5.4.1 主性能对比',
    '5.4.3 数据异构强度分析':'5.4.2 数据异构强度分析',
    '5.4.4 本地训练强度分析':'5.4.3 本地训练强度分析',
    '5.4.5 共享特征池规模分析':'5.6.1 共享特征池规模消融',
    '5.4.6 特征蒸馏充分性分析':'5.6.2 特征蒸馏充分性消融',
    '5.4.7 特征蒸馏机制诊断':'6.1 特征蒸馏机制诊断',
    '5.3 隐私验证':'6.2 隐私验证',
    '5.3.1 隐私验证思路':'6.2.1 隐私验证思路',
    '5.3.2 敏感特征剪裁与噪声设置':'6.2.2 敏感特征剪裁与噪声设置',
    '5.3.3 模型反演攻击验证':'6.2.3 模型反演攻击验证',
    '5.3.4 成员推断攻击验证':'6.2.4 成员推断攻击验证',
    '5.3.5 结果分析':'6.2.5 隐私验证结果分析',
    '5.4.8 跨联邦优化器可插拔验证':'6.3 跨联邦优化器可插拔验证',
    '5.4.9 训练开销分析':'6.4 训练开销分析',
}
for p in doc.paragraphs:
    t=p.text.strip()
    if t in renames:
        # decide level by new text
        new=renames[t]
        level=2 if new.count('.')==1 else 3
        set_text(p,new,level=level)

# Fix chapter titles and normalize heading styles in new chapters.
for p in doc.paragraphs:
    t=p.text.strip()
    if t in ['第5章 实验设计与性能评估','第6章 机制验证与工程分析','结  论']:
        fmt_heading(p,1)
    elif t.startswith('5.') or t.startswith('6.'):
        # table/figure captions also start with 图/表, not here
        parts=t.split(' ')[0]
        if parts.count('.')==1:
            fmt_heading(p,2)
        elif parts.count('.')==2:
            fmt_heading(p,3)
    elif t.startswith('图 ') or t.startswith('表 '):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=None
        p.paragraph_format.line_spacing=Pt(LINE_PT); p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
        for r in p.runs: set_run_font(r)

# Update visible TOC stale chapter entries lightly.
for p in doc.paragraphs:
    t=p.text.strip()
    if t.startswith('第5章 实验设计与结果分析') or t.startswith('第5章 可行性分析与性能参数'):
        set_text(p,'第5章 实验设计与性能评估.............................................................................5')
    elif t.startswith('结  论') and p.style.name.startswith('toc'):
        set_text(p,'结  论..................................................................................................................................7')

# Ask Word to update fields/TOC.
settings=doc.settings._element
upd=settings.find(qn('w:updateFields'))
if upd is None:
    upd=OxmlElement('w:updateFields'); settings.append(upd)
upd.set(qn('w:val'),'true')

doc.save(str(DOC_PATH))
print(DOC_PATH)
print(BACKUP_PATH)

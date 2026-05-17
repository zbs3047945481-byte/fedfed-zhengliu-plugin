from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


DOC = Path("/Users/zhubingshuo/毕设资料/2-复制.docx")


ABSTRACT_CN = (
    "联邦学习能够在原始数据不出本地的条件下协同训练模型，但真实场景中的客户端数据通常呈现显著 Non-IID 分布，"
    "导致本地更新方向不一致、全局收敛变慢和模型精度下降。针对该问题，本文设计并实现了一种即插即用的 FedFed "
    "特征蒸馏插件，将特征蒸馏、共享特征池构建和共享特征辅助训练从具体联邦优化流程中解耦，使其能够在不改变主聚合逻辑的条件下接入不同联邦学习基线。"
    "本文首先构建多基线联邦学习框架，保留 FedAvg、FedProx、SCAFFOLD、FedAvgM 和 FedNova 等主训练流程，并为 FedFed 插件提供独立辅助信息通道。"
    "随后，本文实现基于图像空间特征蒸馏的插件模块：客户端通过生成器和蒸馏分类器学习分类相关敏感特征，服务器按类别维护共享特征池，"
    "并在正式联邦训练阶段向客户端下发共享特征，以补充本地数据缺失的类别信息。"
    "在 CIFAR-10 强异构场景下，FedFed 插件在 FedAvg 基线上将最高测试准确率由 63.89% 提升至 88.95%；"
    "在更强异构设置下，最高准确率由 38.15% 提升至 84.80%；在本地训练轮数增大的场景下，最高准确率由 59.60% 提升至 91.90%。"
    "与 FedFed 原方法相比，本文插件化实现牺牲少量主性能，但换取了更清晰的模块边界和跨算法复用能力：在 FedAvg 主实验、更强异构、本地训练增强和 FedProx 扩展设置下，"
    "插件结果分别比原方法低 3.39、5.22、1.34 和 1.63 个百分点。跨优化器实验表明，该插件能够接入多种联邦优化器；机制诊断与消融实验表明，"
    "特征蒸馏阶段、共享特征参与主训练以及蒸馏增强策略均对性能具有关键作用。隐私实验显示，剪裁与噪声能够降低共享特征的可反演性，但不能稳定抑制成员推断风险。"
    "本文工作验证了 FedFed 特征蒸馏思想的插件化可行性，为异构联邦学习中的辅助信息复用提供了一种实现路径。"
)

KEYWORDS_CN = "关键词：联邦学习；数据异构；特征蒸馏；FedFed 插件；Non-IID；隐私验证"

ABSTRACT_EN = (
    "Federated learning enables collaborative model training without directly sharing raw data, but non-independent and identically "
    "distributed client data often cause client drift, slow convergence, and degraded accuracy. This thesis designs and implements a "
    "plug-and-play FedFed feature distillation plugin. The plugin decouples feature distillation, shared feature pool construction, "
    "and shared-feature-assisted training from specific federated optimizers, allowing the FedFed mechanism to be reused without "
    "changing the main aggregation rule. A multi-baseline framework is built with FedAvg, FedProx, SCAFFOLD, FedAvgM, and FedNova, "
    "and an independent auxiliary channel is introduced for the plugin. The plugin learns classification-related sensitive features "
    "through image-space distillation, maintains a class-wise shared feature pool on the server, and distributes shared features to "
    "clients during formal federated training. On CIFAR-10, the plugin improves FedAvg best accuracy from 63.89% to 88.95% under "
    "strong heterogeneity, from 38.15% to 84.80% under stronger heterogeneity, and from 59.60% to 91.90% with more local training. "
    "Compared with the original FedFed method, the plugin is 3.39, 5.22, 1.34, and 1.63 percentage points lower in the main FedAvg, "
    "stronger heterogeneity, stronger local training, and FedProx extension settings, respectively, while obtaining clearer modular "
    "boundaries and cross-algorithm reusability. Mechanism diagnostics, ablation studies, privacy evaluation, and optimizer extension "
    "experiments show that feature distillation and shared-feature-assisted training are essential to the plugin, and that clipping "
    "with Gaussian noise reduces feature invertibility but does not consistently suppress membership inference risk."
)

KEYWORDS_EN = (
    "Keywords: federated learning; data heterogeneity; feature distillation; FedFed plugin; privacy evaluation"
)

CONCLUSION = [
    "本文围绕异构联邦学习中的性能退化问题，设计并实现了一种即插即用的 FedFed 特征蒸馏插件。该插件不重新定义联邦优化算法，而是在原有主训练流程之外建立独立辅助通道，将特征蒸馏、共享特征池和共享特征辅助训练封装为可启停模块。关闭插件时，系统执行原始联邦基线；开启插件时，客户端获得跨客户端共享的分类相关特征，从而缓解本地类别分布偏斜带来的训练偏移。",
    "本文首先实现了多基线联邦学习框架下的插件化接入机制。系统保留 FedAvg、FedProx、SCAFFOLD、FedAvgM 和 FedNova 的主流程，将插件状态维护、共享特征收集和辅助训练过程从基线算法中分离，提高了方法复用性和实验可比性。",
    "实验结果表明，FedFed 插件能够在强 Non-IID 场景下显著提升分类性能。在 CIFAR-10、α=0.1、E=1 设置下，FedAvg 最高准确率为 63.89%，开启插件后提升至 88.95%；在 α=0.05 的更强异构设置下，插件将最高准确率由 38.15% 提升至 84.80%；在 E=5 设置下，插件将最高准确率由 59.60% 提升至 91.90%。这些结果说明，共享特征辅助训练能够有效缓解数据异构造成的性能下降。",
    "与 FedFed 原方法相比，本文插件化实现以少量核心性能损失换取了更强的系统复用能力。在 FedAvg 主实验、更强异构、本地训练增强和 FedProx 扩展设置下，插件结果分别比原方法低 3.39、5.22、1.34 和 1.63 个百分点。该结果说明，FedFed 机制被封装为即插即用插件后，仍能保留主要性能收益，并使特征蒸馏模块能够独立接入不同联邦主流程。",
    "机制诊断和消融实验验证了插件关键模块的必要性。敏感特征具有较强分类能力，鲁棒特征分类能力较弱，说明特征蒸馏过程能够将分类相关信息集中到共享特征中。关闭特征蒸馏阶段、关闭共享特征主训练或关闭蒸馏增强策略后，模型性能均明显下降，表明插件性能来自多个模块的协同作用。",
    "跨优化器实验表明，FedFed 插件具备跨联邦优化器接入能力。插件在 FedProx、FedAvgM 和 FedNova 上带来明显提升，在 SCAFFOLD 上提升较小。SCAFFOLD 通过控制变量校正客户端梯度方向，FedFed 插件通过共享特征改变本地训练信号，二者同时作用于客户端更新过程，因此插件增益会受到底层优化机制影响。",
    "隐私验证表明，范数剪裁与高斯噪声能够降低共享特征的模型反演风险。最佳配置下，分类代理准确率仅小幅下降，模型反演 PSNR 由 18.87 降至 16.58，说明共享特征中的可恢复图像信息减少。但成员推断实验显示，该机制不能稳定降低 MIA 风险，因此本文方法不属于严格隐私保护机制，仍需结合动态噪声、正则化训练或差分隐私方法进一步改进。",
    "总体来看，本文完成了 FedFed 特征蒸馏思想从算法机制到插件化系统的实现，并通过性能实验、原方法对照、机制诊断、消融实验、隐私验证、跨优化器实验和开销分析验证了其有效性与局限性。后续工作可从三个方向展开：一是优化特征蒸馏与共享池维护策略，降低插件训练开销；二是设计更稳定的共享特征隐私保护机制；三是扩展到更多数据集、模型结构和真实联邦部署场景，进一步验证插件化特征蒸馏的泛化能力。",
]


def has_drawing(paragraph):
    xml = paragraph._p.xml
    return "w:drawing" in xml or "w:pict" in xml


def clear_para(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def paragraph_has_page_break(paragraph):
    return 'w:type="page"' in paragraph._p.xml or "w:lastRenderedPageBreak" in paragraph._p.xml


def ensure_break_before(doc, target_text):
    for idx, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == target_text and idx > 0:
            prev = doc.paragraphs[idx - 1]
            if not paragraph_has_page_break(prev):
                prev.add_run().add_break(WD_BREAK.PAGE)
            return


def set_para_text(paragraph, text):
    if has_drawing(paragraph):
        raise ValueError("refuse to overwrite paragraph containing drawing")
    clear_para(paragraph)
    if text:
        paragraph.add_run(text)


def set_rfonts(run, east="宋体", latin="Times New Roman", size=12, bold=None):
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), east)
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:cs"), latin)
    run.font.name = latin
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def fmt_para(paragraph, *, east="宋体", latin="Times New Roman", size=12, bold=None,
             align=None, first=24, left=None, before=0, after=0, line=22):
    pf = paragraph.paragraph_format
    if align is not None:
        paragraph.alignment = align
    pf.space_before = Pt(before) if before is not None else None
    pf.space_after = Pt(after) if after is not None else None
    if line is not None:
        pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        pf.line_spacing = Pt(line)
    if first is None:
        pf.first_line_indent = None
    else:
        pf.first_line_indent = Pt(first)
    pf.left_indent = Pt(left) if left is not None else None
    for run in paragraph.runs:
        set_rfonts(run, east=east, latin=latin, size=size, bold=bold)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge_data = kwargs.get(edge)
        tag = "w:{}".format(edge)
        element = tcBorders.find(qn(tag))
        if edge_data is None:
            continue
        if element is None:
            element = OxmlElement(tag)
            tcBorders.append(element)
        for key, value in edge_data.items():
            element.set(qn("w:{}".format(key)), str(value))


def no_border():
    return {"val": "nil"}


def single(sz=8):
    return {"val": "single", "sz": sz, "space": "0", "color": "000000"}


def format_three_line_table(table):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top=no_border(), bottom=no_border(), left=no_border(), right=no_border(), insideH=no_border(), insideV=no_border())
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
                p.paragraph_format.line_spacing = Pt(18)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    set_rfonts(r, size=10.5)
    if not table.rows:
        return
    first = table.rows[0]
    last = table.rows[-1]
    for cell in first.cells:
        set_cell_border(cell, top=single(10), bottom=single(6), left=no_border(), right=no_border())
    for cell in last.cells:
        set_cell_border(cell, bottom=single(10), left=no_border(), right=no_border())


def normalize_text_fixes(doc):
    fixes = {
        "数据划分采用 Dirichlet 分布构造标签分布偏斜，并启用客户端样本数量偏斜。Dirichlet 参数 用于控制数据异构强度， 越小，客户端之间的类别分布差异越强。默认设置为 ，对应较强的 Non-IID 场景。":
        "数据划分采用 Dirichlet 分布构造标签分布偏斜，并启用客户端样本数量偏斜。Dirichlet 参数 α 用于控制数据异构强度，α 越小，客户端之间的类别分布差异越强。默认设置为 α=0.1，对应较强的 Non-IID 场景。"
    }
    for p in doc.paragraphs:
        text = " ".join(p.text.split())
        if text in fixes and not has_drawing(p):
            set_para_text(p, fixes[text])


def main():
    doc = Document(DOC)

    # Page and margin baseline used by the current thesis template.
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(3.0)
        section.bottom_margin = Cm(3.0)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.8)

    normalize_text_fixes(doc)

    # Abstracts.
    set_para_text(doc.paragraphs[61], ABSTRACT_CN)
    set_para_text(doc.paragraphs[62], "")
    set_para_text(doc.paragraphs[63], KEYWORDS_CN)
    set_para_text(doc.paragraphs[65], ABSTRACT_EN)
    set_para_text(doc.paragraphs[66], "")
    set_para_text(doc.paragraphs[67], KEYWORDS_EN)
    set_para_text(doc.paragraphs[68], "")

    # Keep the second title page independent from the originality statement page.
    if not paragraph_has_page_break(doc.paragraphs[37]):
        doc.paragraphs[37].add_run().add_break(WD_BREAK.PAGE)
    for heading in (
        "5.4 FedFed 插件主性能验证",
        "5.5.2 FedFed 插件的确定",
        "5.5.3 实验配置收敛过程",
        "5.6.3 结构组件消融实验",
        "6.2.3 模型反演攻击验证",
        "6.4 训练开销分析",
    ):
        ensure_break_before(doc, heading)

    # Conclusion.
    start = next(
        i for i, p in enumerate(doc.paragraphs)
        if re.sub(r"\s+", "", p.text) == "结论" and i > 100
    )
    ref = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "参考文献" and i > start)
    slots = list(range(start + 1, ref))
    for idx, text in zip(slots, CONCLUSION):
        set_para_text(doc.paragraphs[idx], text)
    for idx in slots[len(CONCLUSION):]:
        set_para_text(doc.paragraphs[idx], "")

    # Global paragraph formatting.
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        compact = re.sub(r"\s+", "", text)
        if not text and not has_drawing(p):
            # Keep blank lines compact but compatible with the template.
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            p.paragraph_format.line_spacing = Pt(22)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            continue

        if i <= 59:
            # Cover and declaration pages.
            if text == "本科毕业设计（论文）":
                fmt_para(p, east="黑体", latin="Times New Roman", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=None, line=22)
            elif text in {"燕山大学本科毕业设计（论文）", "学位论文原创性声明", "学位论文版权使用授权书"}:
                fmt_para(p, east="黑体", latin="Times New Roman", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=None, line=22)
            elif text == "一种即插即用的联邦学习特征蒸馏模块设计与实现":
                fmt_para(p, east="黑体", latin="Times New Roman", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=None, line=22)
            elif text.startswith("论文题目"):
                fmt_para(p, east="宋体", latin="Times New Roman", size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, first=None, line=22)
            elif text.startswith(("作者姓名", "专", "指导教师")):
                fmt_para(p, east="宋体", latin="Times New Roman", size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, first=None, left=180, line=22)
            elif re.match(r"^20\d{2}年\d+月$", text):
                fmt_para(p, east="宋体", latin="Times New Roman", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=None, line=22)
            else:
                fmt_para(p, east="宋体", latin="Times New Roman", size=12, bold=False, align=p.alignment, first=24 if len(text) > 25 else None, line=22)
            continue

        if compact in {"摘要", "目录"}:
            fmt_para(p, east="黑体", latin="Times New Roman", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=None, before=22.35, after=17.85, line=22)
        elif text == "Abstract":
            fmt_para(p, east="Times New Roman", latin="Times New Roman", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=None, before=22.35, after=17.85, line=22)
        elif 70 <= i <= 107:
            left = 15 if re.match(r"^\d+\.\d", text) else 0
            fmt_para(p, east="宋体", latin="Times New Roman", size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, first=None, left=left, line=22)
        elif text.startswith("关键词："):
            fmt_para(p, east="宋体", latin="Times New Roman", size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, first=None, line=22)
            if p.runs:
                p.runs[0].bold = False
        elif text.startswith("Keywords:"):
            fmt_para(p, east="Times New Roman", latin="Times New Roman", size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, first=None, line=22)
        elif re.match(r"^第\d+章", text) or compact in {"结论", "参考文献", "致谢"}:
            fmt_para(p, east="黑体", latin="Times New Roman", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=None, before=22.35, after=17.85, line=22)
        elif re.match(r"^\d+\.\d+\.\d+", text):
            fmt_para(p, east="黑体", latin="Times New Roman", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, first=None, before=7.8, after=7.8, line=22)
        elif re.match(r"^\d+\.\d+", text):
            fmt_para(p, east="黑体", latin="Times New Roman", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, first=None, before=7.8, after=7.8, line=22)
        elif re.match(r"^[图表]\s*\d+-\d+", text):
            fmt_para(p, east="宋体", latin="Times New Roman", size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, first=None, before=0, after=6, line=18)
        elif re.search(r"（\d+-\d+）$", text) and len(text) < 120:
            fmt_para(p, east="宋体", latin="Times New Roman", size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, first=None, line=22)
        else:
            fmt_para(p, east="宋体", latin="Times New Roman", size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=24, line=22)

    # Cover information table.
    if doc.tables:
        for row in doc.tables[0].rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
                    p.paragraph_format.line_spacing = Pt(22)
                    for r in p.runs:
                        set_rfonts(r, size=12)

    for table in doc.tables[1:]:
        table.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_three_line_table(table)

    doc.save(DOC)
    print("updated", DOC, "paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "images", len(doc.inline_shapes))


if __name__ == "__main__":
    main()

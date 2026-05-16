from pathlib import Path
from shutil import copy2
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

DOC_PATH = Path('/Users/zhubingshuo/毕设资料/2-复制.docx')
BACKUP_PATH = DOC_PATH.with_name('2-复制.backup-before-5.4.docx')
ROOT = Path('/Users/zhubingshuo/plug-and-play-fedfed-fedavg-plugin')
FIG_DIR = ROOT / 'tmp/docs/figures_54'
FIGS = {
    'main_bar': FIG_DIR / 'fig_5_5_main_bar.png',
    'main_curve': FIG_DIR / 'fig_5_6_main_curve.png',
    'heter_bar': FIG_DIR / 'fig_5_7_heter_bar.png',
    'heter_gain': FIG_DIR / 'fig_5_8_heter_gain.png',
    'local_bar': FIG_DIR / 'fig_5_9_local_bar.png',
    'pool_bar': FIG_DIR / 'fig_5_10_pool_bar.png',
    'distill_bar': FIG_DIR / 'fig_5_11_distill_bar.png',
}

if not BACKUP_PATH.exists():
    copy2(DOC_PATH, BACKUP_PATH)

for key, path in FIGS.items():
    if not path.exists():
        raise FileNotFoundError(path)

doc = Document(str(DOC_PATH))

# Font helpers.
def set_run_font(run, size=10.5, bold=False, italic=False, name='宋体'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def set_para_body(paragraph, first_indent=True):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    if first_indent:
        fmt.first_line_indent = Pt(21)  # about 2 Chinese chars at 10.5 pt
    for r in paragraph.runs:
        set_run_font(r, 10.5)


def set_heading(paragraph, level):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = None
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(6 if level == 2 else 4)
    fmt.space_after = Pt(3)
    fmt.keep_with_next = True
    for r in paragraph.runs:
        set_run_font(r, 14 if level == 2 else 12, bold=True, name='黑体')


def clear_para(p):
    for r in list(p.runs):
        p._p.remove(r._r)


def add_plain_para_before(ref_el, text, style=None, first_indent=True):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    set_run_font(run, 10.5)
    set_para_body(p, first_indent=first_indent)
    ref_el.addprevious(p._p)
    return p


def add_heading_before(ref_el, text, level=2):
    p = doc.add_paragraph(style=f'Heading {level}' if f'Heading {level}' in [s.name for s in doc.styles] else None)
    run = p.add_run(text)
    set_heading(p, level)
    ref_el.addprevious(p._p)
    return p


def add_caption_before(ref_el, text, kind='fig'):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = p.paragraph_format
    fmt.first_line_indent = None
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(3 if kind == 'fig' else 6)
    fmt.space_after = Pt(3)
    fmt.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, 10.5, name='宋体')
    ref_el.addprevious(p._p)
    return p


def add_explain_before(ref_el, text):
    return add_plain_para_before(ref_el, text, first_indent=True)


def add_image_before(ref_el, path, width_inches=5.65):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = p.paragraph_format
    fmt.first_line_indent = None
    fmt.line_spacing = 1.0
    fmt.space_before = Pt(3)
    fmt.space_after = Pt(0)
    fmt.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_inches))
    ref_el.addprevious(p._p)
    return p


def add_page_break_before(ref_el):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)
    ref_el.addprevious(p._p)
    return p


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        tag = 'w:{}'.format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        if edge_data is None:
            element.set(qn('w:val'), 'nil')
            element.set(qn('w:sz'), '0')
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), 'auto')
        else:
            for key, value in edge_data.items():
                element.set(qn(f'w:{key}'), str(value))


def set_table_width(table, width_dxa=8730):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.insert(0, tblW)
    tblW.set(qn('w:type'), 'dxa')
    tblW.set(qn('w:w'), str(width_dxa))
    jc = tblPr.first_child_found_in('w:jc')
    if jc is None:
        jc = OxmlElement('w:jc')
        tblPr.append(jc)
    jc.set(qn('w:val'), 'center')
    tblLayout = tblPr.first_child_found_in('w:tblLayout')
    if tblLayout is None:
        tblLayout = OxmlElement('w:tblLayout')
        tblPr.append(tblLayout)
    tblLayout.set(qn('w:type'), 'fixed')


def set_cell_width(cell, width_dxa):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.tcW
    tcW.set(qn('w:type'), 'dxa')
    tcW.set(qn('w:w'), str(width_dxa))


def add_rich_cell_text(cell, parts, size=9.0, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    for part in parts:
        if isinstance(part, str):
            text, sub = part, False
        else:
            text, sub = part
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold)
        if sub:
            r.font.subscript = True


def add_table_before(ref_el, headers, rows, widths, aligns=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    set_table_width(table, sum(widths))
    if aligns is None:
        aligns = [WD_ALIGN_PARAGRAPH.CENTER] * len(headers)
    # Header
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_width(cell, widths[i])
        set_cell_margins(cell)
        add_rich_cell_text(cell, header if isinstance(header, list) else [header], size=8.8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    # Rows
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_width(cell, widths[i])
            set_cell_margins(cell)
            content = value if isinstance(value, list) else [str(value)]
            add_rich_cell_text(cell, content, size=8.8, bold=False, align=aligns[i])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    # Three-line borders: top, header bottom, bottom.
    line = {'val': 'single', 'sz': '8', 'space': '0', 'color': '000000'}
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_border(cell)
            if row_idx == 0:
                set_cell_border(cell, top=line, bottom=line)
            elif row_idx == len(table.rows) - 1:
                set_cell_border(cell, bottom=line)
            else:
                set_cell_border(cell)
    ref_el.addprevious(table._tbl)
    return table


def parts_symbol(base, sub=None):
    return [base] if sub is None else [base, (sub, True)]

# Update visible TOC item for 5.4 if present.
for p in doc.paragraphs:
    if p.text.strip().startswith('5.4 本章小结'):
        clear_para(p)
        r = p.add_run('5.4 FedFed 插件有效性验证.............................................................................6')
        set_run_font(r, 10.5)
        break

# Locate replacement range.
start_p = None
end_p = None
for p in doc.paragraphs:
    text = p.text.strip()
    if text in {'5.4 实验套件设计', '5.4 FedFed 插件有效性验证'}:
        start_p = p
    if text == '结  论' and start_p is not None:
        end_p = p
        break
if start_p is None or end_p is None:
    raise RuntimeError('Could not locate 5.4 replacement range or conclusion heading.')

body = doc._body._element
children = list(body)
start_idx = children.index(start_p._p)
end_idx = children.index(end_p._p)
for el in children[start_idx:end_idx]:
    body.remove(el)

ref_el = end_p._p

# Section content.
add_heading_before(ref_el, '5.4 FedFed 插件有效性验证', 2)
add_plain_para_before(ref_el, 'FedAvg 在 Non-IID 数据场景下容易受到客户端类别分布偏斜影响。由于不同客户端只持有部分类别或类别比例明显不同，本地模型更新方向可能偏离全局分类目标，最终导致全局模型收敛不稳定或测试精度下降。FedFed 插件的设计动机是在不改变 FedAvg 主训练流程的前提下，为客户端提供额外的跨客户端类别判别信息，从而缓解数据异构造成的性能退化。')
add_plain_para_before(ref_el, '本节围绕 FedFed 插件是否有效这一问题展开实验验证。实验将关闭插件的 FedAvg 作为无插件基线，将开启 FedFed 插件的方法作为插件方法。两组方法保持数据集、客户端数量、采样比例、模型结构、优化器和聚合规则一致，区别仅在于是否开启 FedFed 插件。通过这种设置，可以将性能差异主要归因于插件本身。')

add_heading_before(ref_el, '5.4.1 实验验证思路', 3)
add_plain_para_before(ref_el, '本节实验从五个方面验证 FedFed 插件的作用。第一，进行主性能对比，验证在强异构条件下开启插件是否能够显著提升 FedAvg 性能。第二，改变 Dirichlet 参数 α，分析数据异构程度变化时插件收益是否稳定。第三，增大本地训练轮数，观察客户端本地漂移增强时插件是否仍然有效。第四，改变共享特征池容量，分析共享敏感特征数量对性能的影响。第五，改变特征蒸馏训练量，分析蒸馏阶段是否存在充分性和边际收益问题。')
add_plain_para_before(ref_el, '本文使用以下指标进行评价。BestAcc 表示训练过程中全局模型在测试集上达到的最高准确率，用于衡量方法的最佳性能；FinalAcc 表示训练结束或早停时的测试准确率，用于衡量最终稳定性能；BestRound 表示取得最高准确率的通信轮次；FinalRound 表示训练停止时的通信轮次；Gain 表示插件方法相对于无插件基线的 BestAcc 提升。')

add_heading_before(ref_el, '5.4.2 主性能对比', 3)
add_plain_para_before(ref_el, '主性能对比采用较强数据异构设置。Dirichlet 参数 α=0.1，本地训练轮数 E=1。其中，α 用于控制客户端标签分布偏斜程度，数值越小表示数据异构越强；E 表示每轮通信中客户端在本地数据上训练的 epoch 数。')
add_caption_before(ref_el, '表 5-3 强异构场景下的主性能对比', kind='table')
add_table_before(
    ref_el,
    headers=['方法', 'α', 'E', 'BestAcc', 'FinalAcc', 'BestRound', 'FinalRound', 'Gain'],
    rows=[
        ['无插件基线', '0.1', '1', '63.89%', '53.83%', '187', '222', '-'],
        ['FedFed 插件', '0.1', '1', '88.95%', '88.53%', '165', '167', '+25.06%'],
    ],
    widths=[1900, 720, 620, 1120, 1120, 1150, 1200, 900],
    aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
)
add_explain_before(ref_el, '表中，α 表示 Dirichlet 标签分布参数，E 表示本地训练轮数，BestAcc 表示训练过程最高测试准确率，FinalAcc 表示训练结束时测试准确率，BestRound 表示取得最高准确率的通信轮次，FinalRound 表示实际结束轮次，Gain 表示 FedFed 插件相对于无插件基线的最高准确率提升。由表 5-3 可知，在 α=0.1 的强异构场景下，FedFed 插件将 BestAcc 从 63.89% 提升到 88.95%，将 FinalAcc 从 53.83% 提升到 88.53%，说明开启插件后模型性能和训练稳定性均得到明显改善。')
add_image_before(ref_el, FIGS['main_bar'])
add_caption_before(ref_el, '图 5-5 强异构场景下主性能对比', kind='fig')
add_explain_before(ref_el, '图 5-5 对比了两种方法的最高准确率和最终准确率。无插件基线的最终准确率明显低于最高准确率，说明训练后期存在性能回落；FedFed 插件的 BestAcc 与 FinalAcc 接近，说明插件不仅提升了性能上限，也改善了训练稳定性。')
add_image_before(ref_el, FIGS['main_curve'])
add_caption_before(ref_el, '图 5-6 强异构场景下的收敛曲线', kind='fig')
add_explain_before(ref_el, '图 5-6 展示了训练过程中测试准确率随通信轮次的变化。FedFed 插件在较早轮次达到较高精度，并在后续训练中保持稳定；无插件基线虽然中途达到一定精度，但后期波动和下降更明显。这说明 FedFed 插件能够缓解 Non-IID 条件下 FedAvg 的收敛不稳定问题。')

add_heading_before(ref_el, '5.4.3 数据异构强度分析', 3)
add_plain_para_before(ref_el, '为进一步分析插件收益与数据异构程度之间的关系，本文设置 α=0.3、α=0.1 和 α=0.05 三种数据划分。α=0.3 表示中等异构，α=0.1 表示较强异构，α=0.05 表示更强的类别分布偏斜。其余训练参数保持一致。')
add_page_break_before(ref_el)
add_caption_before(ref_el, '表 5-4 不同数据异构强度下的性能对比', kind='table')
add_table_before(
    ref_el,
    headers=['α', 'E', '方法', 'BestAcc', 'FinalAcc', 'Gain'],
    rows=[
        ['0.3', '1', '无插件基线', '71.95%', '67.81%', '-'],
        ['0.3', '1', 'FedFed 插件', '91.98%', '91.14%', '+20.03%'],
        ['0.1', '1', '无插件基线', '63.89%', '53.83%', '-'],
        ['0.1', '1', 'FedFed 插件', '88.95%', '88.53%', '+25.06%'],
        ['0.05', '1', '无插件基线', '38.15%', '34.93%', '-'],
        ['0.05', '1', 'FedFed 插件', '84.80%', '78.64%', '+46.65%'],
    ],
    widths=[750, 650, 2100, 1400, 1400, 1200],
    aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
)
add_explain_before(ref_el, '表中，α 表示数据异构强度控制参数，E 表示本地训练轮数，Gain 表示相同 α 和 E 下 FedFed 插件相对于无插件基线的 BestAcc 提升。由表 5-4 可知，随着 α 变小，无插件基线性能快速下降；相比之下，FedFed 插件在三种异构强度下均保持较高准确率，并在 α=0.05 时取得 46.65 个百分点的最大提升。')
add_image_before(ref_el, FIGS['heter_bar'])
add_caption_before(ref_el, '图 5-7 不同数据异构强度下的性能对比', kind='fig')
add_explain_before(ref_el, '图 5-7 显示，FedFed 插件在不同 α 下均明显优于无插件基线。无插件基线对 α 的变化更加敏感，而 FedFed 插件的准确率下降幅度较小，说明共享敏感特征能够提高模型对数据异构的鲁棒性。')
add_image_before(ref_el, FIGS['heter_gain'], width_inches=5.25)
add_caption_before(ref_el, '图 5-8 不同数据异构强度下的插件增益', kind='fig')
add_explain_before(ref_el, '图 5-8 展示了 FedFed 插件在不同异构强度下的准确率增益。随着 α 降低，插件增益整体增大，说明客户端数据越偏斜，共享敏感特征提供的跨客户端补充信息越重要。')

add_heading_before(ref_el, '5.4.4 本地训练强度分析', 3)
add_plain_para_before(ref_el, '本地训练轮数增大时，客户端会在本地数据上执行更多更新。该设置能够减少通信次数，但在 Non-IID 场景下也可能加剧客户端更新方向差异，形成更明显的 client drift。为验证 FedFed 插件是否能够缓解这一问题，本文将本地训练轮数设置为 E=5，并与无插件基线进行对比。')
add_page_break_before(ref_el)
add_caption_before(ref_el, '表 5-5 本地训练强度增大时的性能对比', kind='table')
add_table_before(
    ref_el,
    headers=['α', 'E', '方法', 'BestAcc', 'FinalAcc', 'BestRound', 'FinalRound', 'Gain'],
    rows=[
        ['0.1', '5', '无插件基线', '59.60%', '58.57%', '58', '120', '-'],
        ['0.1', '5', 'FedFed 插件', '91.90%', '91.01%', '132', '167', '+32.30%'],
    ],
    widths=[650, 560, 1700, 1080, 1080, 1120, 1180, 900],
    aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
)
add_explain_before(ref_el, '表中，E 表示每轮通信中的本地训练 epoch 数。较大的 E 会增强本地训练强度，也更容易放大 Non-IID 数据下的客户端漂移。由表 5-5 可知，当 E=5 时，无插件基线最高准确率为 59.60%，而 FedFed 插件达到 91.90%，提升 32.30 个百分点，说明插件能够在本地漂移增强时继续保持有效。')
add_image_before(ref_el, FIGS['local_bar'])
add_caption_before(ref_el, '图 5-9 本地训练强度增大时的性能对比', kind='fig')
add_explain_before(ref_el, '图 5-9 显示，在 E=5 的设置下，FedFed 插件的最高准确率和最终准确率均明显高于无插件基线。这说明共享敏感特征不仅能够补充类别信息，还能在一定程度上约束客户端本地训练方向，使本地更新更接近全局任务目标。')

add_heading_before(ref_el, '5.4.5 共享特征池规模分析', 3)
add_plain_para_before(ref_el, '共享特征池用于保存各客户端上传的性能敏感特征，并在正式训练阶段下发给客户端辅助训练。共享池容量过小可能导致类别覆盖不足，容量较大则能够保留更多跨客户端判别信息。为分析共享池规模对插件性能的影响，本文设置低容量共享池、中等容量共享池和不设额外容量限制三种配置。')
add_caption_before(ref_el, '表 5-6 不同共享特征池规模下的性能对比', kind='table')
add_table_before(
    ref_el,
    headers=['共享池设置', '上传总量上限', '单类上传上限', '池总量上限', '单类池容量', 'BestAcc', 'FinalAcc'],
    rows=[
        ['低容量共享池', '100', '4', '800', '80', '80.31%', '75.30%'],
        ['中等容量共享池', '1000', '20', '4000', '400', '80.12%', '75.48%'],
        ['不设额外容量限制', '不限制', '不限制', '不限制', '不限制', '90.33%', '89.50%'],
    ],
    widths=[1700, 1250, 1200, 1150, 1050, 1050, 1050],
    aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
)
add_explain_before(ref_el, '表中，上传总量上限表示单个客户端最多上传的共享特征数量，单类上传上限表示每个类别最多上传的特征数量，池总量上限表示服务器共享池保存的最大特征数量，单类池容量表示服务器对每一类共享特征的容量限制。由表 5-6 可知，不设额外容量限制时，FedFed 插件最高准确率达到 90.33%，明显高于低容量和中等容量设置，说明共享特征池容量是影响插件效果的重要因素。')
add_image_before(ref_el, FIGS['pool_bar'])
add_caption_before(ref_el, '图 5-10 不同共享特征池规模下的性能对比', kind='fig')
add_explain_before(ref_el, '图 5-10 显示，不设额外容量限制的共享池在 BestAcc 和 FinalAcc 上均明显优于受限共享池。说明在强异构条件下，FedFed 插件需要足够丰富的共享敏感特征来覆盖客户端缺失类别，过小的共享池会削弱插件作用。')

add_heading_before(ref_el, '5.4.6 特征蒸馏充分性分析', 3)
add_plain_para_before(ref_el, '特征蒸馏阶段决定共享敏感特征的质量。若蒸馏不足，提取出的敏感特征可能缺少稳定判别信息；若蒸馏训练量过大，则可能增加训练开销，并不一定继续提高正式训练性能。本文通过改变蒸馏通信轮数 Rd 和本地蒸馏轮数 Ed 分析蒸馏充分性。')
add_plain_para_before(ref_el, '需要说明的是，本组实验中特征蒸馏损失权重 λfd 保持为 2.0，重构损失权重 λrec 保持为 5.0，原图分类损失权重 λx 保持为 0.4。因此，本组实验主要分析蒸馏训练量变化，而不是蒸馏损失权重变化。')
add_caption_before(ref_el, '表 5-7 不同特征蒸馏训练量下的性能对比', kind='table')
add_table_before(
    ref_el,
    headers=['蒸馏设置', parts_symbol('λ', 'fd'), parts_symbol('R', 'd'), parts_symbol('E', 'd'), 'BestAcc', 'FinalAcc', 'BestRound'],
    rows=[
        ['短程蒸馏', '2.0', '5', '1', '89.73%', '89.43%', '142'],
        ['标准蒸馏', '2.0', '15', '1', '88.41%', '87.50%', '117'],
        ['长程蒸馏', '2.0', '30', '1', '89.26%', '88.73%', '157'],
        ['长程增强蒸馏', '2.0', '30', '2', '90.30%', '88.19%', '212'],
    ],
    widths=[1800, 820, 720, 720, 1200, 1200, 1200],
    aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
)
add_explain_before(ref_el, '表中，λfd 表示特征蒸馏分类损失权重，Rd 表示特征蒸馏阶段的通信轮数，Ed 表示每轮蒸馏中的本地训练 epoch 数。Rd 和 Ed 越大，表示蒸馏阶段训练量越大。由表 5-7 可知，不同蒸馏训练量下插件均能取得较高准确率，但性能并不随 Rd 或 Ed 增大而单调提升。')
add_image_before(ref_el, FIGS['distill_bar'])
add_caption_before(ref_el, '图 5-11 不同特征蒸馏训练量下的性能对比', kind='fig')
add_explain_before(ref_el, '图 5-11 显示，不同蒸馏训练量下的性能差异相对有限，且不存在明显单调趋势。这说明蒸馏阶段的主要作用是获得可用的敏感特征表示；当特征分离质量达到一定程度后，继续增加蒸馏训练量的边际收益有限。')

add_heading_before(ref_el, '5.4.7 结果分析', 3)
for text in [
    '综合上述实验，可以得到以下结论。第一，FedFed 插件能够显著提升 FedAvg 在 Non-IID 图像分类任务中的性能。在 α=0.1、E=1 的强异构设置下，FedFed 插件将最高准确率从 63.89% 提升到 88.95%，说明插件提供的共享敏感特征能够有效补充客户端本地数据中的类别信息缺口。',
    '第二，FedFed 插件在更强数据异构场景下收益更明显。当 α=0.05 时，无插件基线最高准确率下降到 38.15%，而 FedFed 插件仍达到 84.80%，提升 46.65 个百分点。这表明插件主要针对 Non-IID 数据分布带来的客户端类别偏斜问题发挥作用。',
    '第三，FedFed 插件能够缓解本地训练增强带来的 client drift。当本地训练轮数增大到 E=5 时，无插件基线性能下降，而 FedFed 插件仍保持 91.90% 的最高准确率，说明共享敏感特征能够使客户端本地训练目标更接近全局任务。',
    '第四，共享特征池容量对插件效果影响较大。不设额外容量限制时，插件性能明显优于低容量和中等容量共享池，说明强异构场景下需要足够丰富的共享敏感特征覆盖不同类别。',
    '第五，特征蒸馏训练量存在边际收益。固定 λfd=2.0 时，增加 Rd 或 Ed 并不稳定提升最终准确率。该结果说明后续优化应更关注共享特征质量、类别覆盖和共享池筛选策略，而不是单纯增加蒸馏阶段训练量。',
    '综上，FedFed 插件在保持 FedAvg 主流程不变的情况下，通过开启插件机制显著提升了强异构联邦图像分类性能。实验结果验证了本文插件化设计的有效性，也说明 FedFed 特征蒸馏思想可以以即插即用方式接入 FedAvg 框架。',
]:
    add_plain_para_before(ref_el, text)

# Apply subscript formatting in paragraphs for specific symbols.
def subscript_terms(paragraph):
    text = paragraph.text
    replacements = [('λfd', ('λ', 'fd')), ('λrec', ('λ', 'rec')), ('λx', ('λ', 'x')), ('Rd', ('R', 'd')), ('Ed', ('E', 'd'))]
    if not any(a in text for a, _ in replacements):
        return
    clear_para(paragraph)
    i = 0
    while i < len(text):
        matched = None
        for raw, parts in replacements:
            if text.startswith(raw, i):
                matched = (raw, parts)
                break
        if matched:
            raw, (base, sub) = matched
            r = paragraph.add_run(base)
            set_run_font(r, 10.5)
            r2 = paragraph.add_run(sub)
            set_run_font(r2, 10.5)
            r2.font.subscript = True
            i += len(raw)
        else:
            r = paragraph.add_run(text[i])
            set_run_font(r, 10.5)
            i += 1
    set_para_body(paragraph, first_indent=True)

for p in doc.paragraphs:
    subscript_terms(p)

# Encourage Word to update fields/TOC when opened.
settings = doc.settings._element
update = settings.find(qn('w:updateFields'))
if update is None:
    update = OxmlElement('w:updateFields')
    settings.append(update)
update.set(qn('w:val'), 'true')

# Save in place.
doc.save(str(DOC_PATH))
print(f'saved={DOC_PATH}')
print(f'backup={BACKUP_PATH}')

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
DOC='/Users/zhubingshuo/毕设资料/2-复制.docx'
d=Document(DOC)
section=None
rules={
 '5.4.1 主性能对比': {'表 6-2':'表 5-1','图 6-2':'图 5-1','图 6-3':'图 5-2'},
 '5.4.2 数据异构强度分析': {'表 6-3':'表 5-2','图 6-4':'图 5-3','图 6-5':'图 5-4'},
 '5.4.3 本地训练强度分析': {'表 6-2':'表 5-3','图 6-2':'图 5-5'},
 '5.6.1 共享特征池规模消融': {'表 6-3':'表 5-4','图 6-3':'图 5-6'},
 '5.6.2 特征蒸馏充分性消融': {'表 6-2':'表 5-5','图 6-4':'图 5-7'},
}
for p in d.paragraphs:
    t=p.text.strip()
    if t in rules:
        section=t
        continue
    if t.startswith('5.') and t not in rules and p.style.name.startswith('Heading'):
        section=None
    if t.startswith('第6章'):
        section=None
    repl=rules.get(section)
    if not repl:
        continue
    for r in p.runs:
        txt=r.text
        for old,new in repl.items(): txt=txt.replace(old,new)
        r.text=txt
    # fallback if split runs remain
    txt=p.text
    if any(old in txt for old in repl):
        for old,new in repl.items(): txt=txt.replace(old,new)
        for r in list(p.runs): p._p.remove(r._r)
        r=p.add_run(txt)
        r.font.name='宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(12)
    if p.text.strip().startswith('图 ') or p.text.strip().startswith('表 '):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=None
        p.paragraph_format.line_spacing=Pt(22); p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
        for r in p.runs:
            r.font.name='宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(12)

d.save(DOC)

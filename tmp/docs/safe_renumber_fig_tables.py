from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
DOC='/Users/zhubingshuo/毕设资料/2-复制.docx'
d=Document(DOC)
mapping={
'表 5-10':'表 6-5','图 5-14':'图 6-7','图 5-13':'图 6-6','图 5-12':'图 6-1','表 5-9':'表 6-4','表 5-8':'表 6-1',
'图 5-11':'图 5-7','图 5-10':'图 5-6','表 5-7':'表 5-5','表 5-6':'表 5-4','图 5-9':'图 5-5','表 5-5':'表 5-3','图 5-8':'图 5-4','图 5-7':'图 5-3','表 5-4':'表 5-2','图 5-6':'图 5-2','图 5-5':'图 5-1','表 5-3':'表 5-1',
'图 5-4':'图 6-5','表 5-2':'表 6-3','图 5-3':'图 6-4','图 5-2':'图 6-3','表 5-1':'表 6-2','图 5-1':'图 6-2',
}
items=sorted(mapping.items(), key=lambda kv: len(kv[0]), reverse=True)
for p in d.paragraphs:
    full=p.text
    if not any(old in full for old,_ in items):
        continue
    has_drawing='w:drawing' in p._p.xml or 'w:pict' in p._p.xml
    changed=False
    for r in p.runs:
        if 'w:drawing' in r._r.xml or 'w:pict' in r._r.xml:
            continue
        txt=r.text
        new=txt
        for old,val in items:
            new=new.replace(old,val)
        if new != txt:
            r.text=new; changed=True
    if not changed and not has_drawing:
        txt=full
        for old,val in items: txt=txt.replace(old,val)
        for r in list(p.runs): p._p.remove(r._r)
        r=p.add_run(txt); r.font.name='宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(12)
    if p.text.strip().startswith(('图 ','表 ')):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=None
        p.paragraph_format.line_spacing=Pt(22); p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
        p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(3)
        for r in p.runs:
            if 'w:drawing' in r._r.xml or 'w:pict' in r._r.xml: continue
            r.font.name='宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(12)
d.save(DOC)
